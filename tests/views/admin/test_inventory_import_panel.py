"""Tests del InventoryImportPanel (Prompt 5b — split de Dialog en Panel + wrapper)."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import Workbook
from PySide6.QtWidgets import QApplication

from collections_app.app_context import AppContext, create_app_context
from collections_app.core.models.aggregates.inventory_import_report import (
    InventoryImportError,
    InventoryImportReport,
    InventoryImportWarning,
)
from collections_app.core.models.card import Card
from collections_app.core.models.code_header import CodeHeader
from collections_app.core.models.code_line import CodeLine
from collections_app.core.models.collection import Collection
from collections_app.views.admin.inventory_import_dialog import (
    InventoryImportPanel,
    _slugify_for_filename,
)

pytestmark = pytest.mark.gui


@pytest.fixture(autouse=True)
def _qapp() -> Iterator[QApplication | None]:
    app = QApplication.instance() or QApplication([])
    yield app  # type: ignore[misc]


@pytest.fixture
def ctx_with_collection() -> Iterator[tuple[AppContext, Collection]]:
    ctx = create_app_context(":memory:")
    try:
        h = ctx.code_headers.create(
            CodeHeader(code_header_id=None, code_header_name="WC", code_max_length=3)
        )
        assert h.code_header_id is not None
        ctx.code_lines.upsert(
            CodeLine(
                code_line_id=None,
                code_header_id=h.code_header_id,
                code_id="ARG",
                code_name="Argentina",
                code_order=1,
            )
        )
        coll = ctx.collections.create(
            Collection(
                collection_id=None,
                collection_name="Mundial 2026",
                card_count=1,
                requires_code=True,
                code_field_name="País",
                code_header_id=h.code_header_id,
            )
        )
        assert coll.collection_id is not None
        ctx.cards.create(
            Card(
                card_id=None,
                collection_id=coll.collection_id,
                code_id="ARG",
                card_number=1,
                card_name="Messi",
            )
        )
        ctx.conn.commit()
        yield ctx, coll
    finally:
        ctx.close()


def test_construct_panel_with_collection(
    qtbot,  # type: ignore[no-untyped-def]
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    ctx, coll = ctx_with_collection
    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    assert panel._mode_replace.isChecked() is True
    assert panel._import_btn.isEnabled() is False


def test_import_button_disabled_without_file(
    qtbot,  # type: ignore[no-untyped-def]
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    ctx, coll = ctx_with_collection
    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    assert panel._import_btn.isEnabled() is False
    panel._file_edit.setText("/tmp/foo.xlsx")
    assert panel._import_btn.isEnabled() is True


def test_download_template_writes_to_downloads(
    qtbot,  # type: ignore[no-untyped-def]
    tmp_path: Path,
    ctx_with_collection: tuple[AppContext, Collection],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Sin diálogo de "dónde guardar": el modelo va auto a Descargas."""
    ctx, coll = ctx_with_collection
    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)

    fake_downloads = tmp_path / "Downloads"
    fake_downloads.mkdir()
    monkeypatch.setattr(
        "collections_app.views.admin.inventory_import_dialog.get_downloads_dir",
        lambda: fake_downloads,
    )
    # Evitar que el success dialog bloquee.
    monkeypatch.setattr(
        "collections_app.views.admin.inventory_import_dialog.ReportSavedDialog.exec",
        lambda self: 0,
    )

    panel._on_download_template()
    files = list(fake_downloads.glob("*.xlsx"))
    assert len(files) == 1
    assert files[0].name.startswith("inventario_")


def test_import_invokes_service_and_emits_signal(
    qtbot,  # type: ignore[no-untyped-def]
    tmp_path: Path,
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    """Click Importar → service y emit del signal con el report."""
    ctx, coll = ctx_with_collection
    file_path = tmp_path / "inv.xlsx"
    wb = Workbook()
    s = wb.active
    assert s is not None
    s.title = "Inventario"
    s.append(["código", "número", "cantidad"])
    s.append(["ARG", 1, 4])
    wb.save(str(file_path))

    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    panel._file_edit.setText(str(file_path))

    captured: list[InventoryImportReport] = []
    panel.import_completed.connect(lambda r: captured.append(r))
    panel._on_import()

    assert len(captured) == 1
    assert captured[0].rows_total == 1
    assert captured[0].rows_applied == 1


def test_renders_report_in_table(
    qtbot,  # type: ignore[no-untyped-def]
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    """Reporte con 1 error + 1 warning → tabla con 2 filas."""
    ctx, coll = ctx_with_collection
    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)

    fake_report = InventoryImportReport(
        rows_total=2,
        rows_applied=0,
        rows_skipped=2,
        errors=[InventoryImportError(row_index=1, message="código vacío")],
        warnings=[
            InventoryImportWarning(
                row_index=2,
                code_id="ARG",
                card_number=1,
                message="cantidad = 0; fila ignorada",
            )
        ],
    )
    panel._render_result(fake_report)
    assert panel._result_table.isHidden() is False
    assert panel._result_table.rowCount() == 2


def test_reset_clears_file_and_results(
    qtbot,  # type: ignore[no-untyped-def]
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    """`reset()` deja el file picker vacío y oculta el bloque de resultado."""
    ctx, coll = ctx_with_collection
    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    panel._file_edit.setText("/tmp/foo.xlsx")
    panel._render_result(InventoryImportReport(rows_total=0, rows_applied=0, rows_skipped=0))
    panel.reset()
    assert panel._file_edit.text() == ""
    assert panel._result_label.isHidden() is True
    assert panel._result_table.rowCount() == 0


def test_set_active_collection_propagates_and_resets(
    qtbot,  # type: ignore[no-untyped-def]
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    """Cambiar de collection actualiza el label y resetea el estado."""
    ctx, coll = ctx_with_collection
    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    panel._file_edit.setText("/tmp/foo.xlsx")

    other = ctx.collections.create(
        Collection(
            collection_id=None,
            collection_name="Otra",
            card_count=0,
            requires_code=False,
            code_field_name=None,
            code_header_id=coll.code_header_id,
        )
    )
    panel.set_active_collection(other)
    assert panel._collection.collection_name == "Otra"
    assert panel._file_edit.text() == ""


def test_slugify_for_filename_strips_accents() -> None:
    assert _slugify_for_filename("Mundial 2026") == "Mundial_2026"
    assert _slugify_for_filename("Mi Colección") == "Mi_Colecci_n"
    assert _slugify_for_filename("!!!") == "inventario"


# ---------------------------------------------------------------------
# Botón "Importar como faltantes"
# ---------------------------------------------------------------------


def test_missing_button_present_and_disabled_without_file(
    qtbot,  # type: ignore[no-untyped-def]
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    """El segundo botón existe y refleja el estado del file picker."""
    ctx, coll = ctx_with_collection
    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    assert panel._import_missing_btn.isEnabled() is False
    panel._file_edit.setText("/tmp/foo.xlsx")
    assert panel._import_missing_btn.isEnabled() is True


def test_missing_button_invokes_import_missing_and_emits_signal(
    qtbot,  # type: ignore[no-untyped-def]
    tmp_path: Path,
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    """Click "Importar como faltantes" llama al service.import_missing
    y emite el signal con el report. La fixture tiene 1 card (ARG-1)
    en la colección; si el archivo lista ARG-1 como faltante, ninguna
    queda como tenida (rows_applied=0)."""
    ctx, coll = ctx_with_collection
    file_path = tmp_path / "faltantes.xlsx"
    wb = Workbook()
    s = wb.active
    assert s is not None
    s.title = "Inventario"
    s.append(["código", "número", "cantidad"])
    s.append(["ARG", 1, 1])
    wb.save(str(file_path))

    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    panel._file_edit.setText(str(file_path))

    captured: list[InventoryImportReport] = []
    panel.import_completed.connect(lambda r: captured.append(r))
    panel._on_import_missing()

    assert len(captured) == 1
    # 1 fila en el Excel (ARG-1 como faltante), 0 cards quedan owned.
    assert captured[0].rows_total == 1
    assert captured[0].rows_applied == 0


def test_missing_button_marks_unlisted_cards_owned(
    qtbot,  # type: ignore[no-untyped-def]
    tmp_path: Path,
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    """Card NO listada como faltante queda con qty=1."""
    ctx, coll = ctx_with_collection
    # Agregar una segunda card para distinguir owned vs missing.
    assert coll.collection_id is not None
    ctx.cards.create(
        Card(
            card_id=None,
            collection_id=coll.collection_id,
            code_id="ARG",
            card_number=2,
            card_name="Di María",
        )
    )
    ctx.conn.commit()

    file_path = tmp_path / "faltantes.xlsx"
    wb = Workbook()
    s = wb.active
    assert s is not None
    s.title = "Inventario"
    s.append(["código", "número", "cantidad"])
    s.append(["ARG", 1, 1])  # ARG-1 falta; ARG-2 NO falta
    wb.save(str(file_path))

    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    panel._file_edit.setText(str(file_path))
    panel._on_import_missing()

    inv_arg1 = ctx.inventory.get_inventory_for_card(coll.collection_id, "ARG", 1)
    inv_arg2 = ctx.inventory.get_inventory_for_card(coll.collection_id, "ARG", 2)
    assert (inv_arg1.quantity if inv_arg1 else 0) == 0
    assert inv_arg2 is not None and inv_arg2.quantity == 1


def test_missing_button_empty_file_asks_confirmation_yes_proceeds(
    qtbot,  # type: ignore[no-untyped-def]
    tmp_path: Path,
    ctx_with_collection: tuple[AppContext, Collection],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Archivo sin filas + confirma Yes → carga todas las cards con qty=1."""
    ctx, coll = ctx_with_collection
    file_path = tmp_path / "vacio.xlsx"
    wb = Workbook()
    s = wb.active
    assert s is not None
    s.title = "Inventario"
    s.append(["código", "número", "cantidad"])
    wb.save(str(file_path))

    from PySide6.QtWidgets import QMessageBox

    monkeypatch.setattr(
        "collections_app.views.admin.inventory_import_dialog.QMessageBox.question",
        lambda *a, **k: QMessageBox.StandardButton.Yes,
    )

    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    panel._file_edit.setText(str(file_path))
    panel._on_import_missing()

    assert coll.collection_id is not None
    inv = ctx.inventory.get_inventory_for_card(coll.collection_id, "ARG", 1)
    assert inv is not None and inv.quantity == 1


def test_missing_button_empty_file_asks_confirmation_no_aborts(
    qtbot,  # type: ignore[no-untyped-def]
    tmp_path: Path,
    ctx_with_collection: tuple[AppContext, Collection],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Archivo sin filas + confirma No → no toca el inventario."""
    ctx, coll = ctx_with_collection
    file_path = tmp_path / "vacio.xlsx"
    wb = Workbook()
    s = wb.active
    assert s is not None
    s.title = "Inventario"
    s.append(["código", "número", "cantidad"])
    wb.save(str(file_path))

    from PySide6.QtWidgets import QMessageBox

    monkeypatch.setattr(
        "collections_app.views.admin.inventory_import_dialog.QMessageBox.question",
        lambda *a, **k: QMessageBox.StandardButton.No,
    )

    captured: list[InventoryImportReport] = []
    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    panel.import_completed.connect(lambda r: captured.append(r))
    panel._file_edit.setText(str(file_path))
    panel._on_import_missing()

    # Sin emission del signal (abortó).
    assert captured == []
    # Inventario intacto.
    assert coll.collection_id is not None
    inv = ctx.inventory.get_inventory_for_card(coll.collection_id, "ARG", 1)
    assert inv is None or inv.quantity == 0


def test_missing_render_shows_owned_and_missing_counts(
    qtbot,  # type: ignore[no-untyped-def]
    ctx_with_collection: tuple[AppContext, Collection],
) -> None:
    """El resumen en modo missing menciona "tenidas" y "faltantes"."""
    ctx, coll = ctx_with_collection
    panel = InventoryImportPanel(ctx=ctx, collection=coll)
    qtbot.addWidget(panel)
    fake_report = InventoryImportReport(
        rows_total=2,
        rows_applied=5,  # cards owned
        rows_skipped=0,
        errors=[],
        warnings=[],
    )
    panel._render_result(fake_report, mode="missing")
    text = panel._result_label.text()
    assert "tenidas" in text.lower()
    assert "faltantes" in text.lower()
