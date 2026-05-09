"""Tests del InventoryImportService — generación de modelos + import."""

from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from collections_app.core.models.card import Card
from collections_app.core.models.code_header import CodeHeader
from collections_app.core.models.code_line import CodeLine
from collections_app.core.models.collection import Collection
from collections_app.core.repositories.cards_repo import CardsRepository
from collections_app.core.repositories.code_headers_repo import CodeHeadersRepository
from collections_app.core.repositories.code_lines_repo import CodeLinesRepository
from collections_app.core.repositories.collections_repo import CollectionsRepository
from collections_app.core.repositories.inventory_repo import InventoryRepository
from collections_app.services.exceptions import ServiceError
from collections_app.services.inventory_import_service import (
    InventoryImportService,
)


@pytest.fixture
def service(db_conn: sqlite3.Connection) -> InventoryImportService:
    return InventoryImportService(db_conn)


@pytest.fixture
def collection_with_codes(db_conn: sqlite3.Connection) -> Collection:
    """Colección Mundial-style: header WC con 3 codes (ARG, BRA, FRA)."""
    headers_repo = CodeHeadersRepository(db_conn)
    lines_repo = CodeLinesRepository(db_conn)
    collections_repo = CollectionsRepository(db_conn)

    header = headers_repo.create(
        CodeHeader(code_header_id=None, code_header_name="WC", code_max_length=3)
    )
    assert header.code_header_id is not None
    for i, (code, name) in enumerate([("ARG", "Argentina"), ("BRA", "Brasil"), ("FRA", "Francia")]):
        lines_repo.upsert(
            CodeLine(
                code_line_id=None,
                code_header_id=header.code_header_id,
                code_id=code,
                code_name=name,
                code_order=i + 1,
            )
        )
    coll = collections_repo.create(
        Collection(
            collection_id=None,
            collection_name="Mundial",
            card_count=15,
            requires_code=True,
            code_field_name="País",
            code_header_id=header.code_header_id,
        )
    )
    db_conn.commit()
    return coll


@pytest.fixture
def collection_without_codes(db_conn: sqlite3.Connection) -> Collection:
    """Colección sin requires_code (solo número + cantidad)."""
    headers_repo = CodeHeadersRepository(db_conn)
    collections_repo = CollectionsRepository(db_conn)
    header = headers_repo.create(
        CodeHeader(code_header_id=None, code_header_name="Empty", code_max_length=3)
    )
    assert header.code_header_id is not None
    coll = collections_repo.create(
        Collection(
            collection_id=None,
            collection_name="Simple",
            card_count=10,
            requires_code=False,
            code_field_name=None,
            code_header_id=header.code_header_id,
        )
    )
    db_conn.commit()
    return coll


# ---------------------------------------------------------------------
# generate_template
# ---------------------------------------------------------------------


def test_generate_template_with_codes_creates_two_sheets(
    tmp_path: Path,
    service: InventoryImportService,
    collection_with_codes: Collection,
) -> None:
    """requires_code=True → 2 pestañas, headers correctos, codes pre-llenados."""
    assert collection_with_codes.collection_id is not None
    dest = tmp_path / "modelo.xlsx"
    result = service.generate_template(
        collection_id=collection_with_codes.collection_id, dest_path=dest
    )
    assert result == dest
    assert dest.exists()

    wb = load_workbook(dest)
    assert wb.sheetnames == ["Inventario", "Códigos"]

    inv = wb["Inventario"]
    assert [c.value for c in inv[1]] == ["código", "número", "cantidad"]
    # Solo headers, sin filas de datos.
    assert inv.max_row == 1

    codes = wb["Códigos"]
    assert [c.value for c in codes[1]] == ["code_id", "code_name"]
    # 3 codes (ARG, BRA, FRA) + header.
    assert codes.max_row == 4
    code_ids = {codes.cell(row=r, column=1).value for r in range(2, 5)}
    assert code_ids == {"ARG", "BRA", "FRA"}


def test_generate_template_without_codes_creates_one_sheet(
    tmp_path: Path,
    service: InventoryImportService,
    collection_without_codes: Collection,
) -> None:
    """requires_code=False → solo pestaña Inventario, sin columna código."""
    assert collection_without_codes.collection_id is not None
    dest = tmp_path / "modelo_simple.xlsx"
    service.generate_template(collection_id=collection_without_codes.collection_id, dest_path=dest)

    wb = load_workbook(dest)
    assert wb.sheetnames == ["Inventario"]
    inv = wb["Inventario"]
    assert [c.value for c in inv[1]] == ["número", "cantidad"]
    assert inv.max_row == 1


def test_generate_template_unknown_collection_raises(
    tmp_path: Path,
    service: InventoryImportService,
) -> None:
    """Collection inexistente → ServiceError (no genera archivo)."""
    dest = tmp_path / "huerfano.xlsx"
    with pytest.raises(ServiceError, match="no existe"):
        service.generate_template(collection_id=999, dest_path=dest)
    assert not dest.exists()


# ---------------------------------------------------------------------
# Helpers para construir archivos de import
# ---------------------------------------------------------------------


def _seed_cards(
    db_conn: sqlite3.Connection, collection: Collection, items: list[tuple[str, int, str]]
) -> dict[tuple[str, int], int]:
    """Crea cards en la DB y devuelve mapping (code_id, num) -> card_id."""
    repo = CardsRepository(db_conn)
    assert collection.collection_id is not None
    out: dict[tuple[str, int], int] = {}
    for code, num, name in items:
        card = repo.create(
            Card(
                card_id=None,
                collection_id=collection.collection_id,
                code_id=code,
                card_number=num,
                card_name=name,
            )
        )
        assert card.card_id is not None
        out[(code, num)] = card.card_id
    db_conn.commit()
    return out


def _write_xlsx(path: Path, header: list[str], rows: list[list[object]]) -> None:
    wb = Workbook()
    sheet = wb.active
    assert sheet is not None
    sheet.title = "Inventario"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    wb.save(str(path))


def _write_csv(path: Path, header: list[str], rows: list[list[object]]) -> None:
    lines = [",".join(header)]
    for row in rows:
        lines.append(",".join(str(c) for c in row))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _qty_for(db_conn: sqlite3.Connection, card_id: int) -> int:
    inv = InventoryRepository(db_conn).get_by_card_id(card_id)
    return inv.quantity if inv is not None else 0


# ---------------------------------------------------------------------
# import_inventory — modo replace
# ---------------------------------------------------------------------


def test_import_xlsx_replace_mode_applies_correctly(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """3 filas válidas → inventory con cantidades exactas."""
    cards_by_key = _seed_cards(
        db_conn,
        collection_with_codes,
        [("ARG", 1, "Messi"), ("BRA", 1, "Neymar"), ("FRA", 1, "Mbappé")],
    )
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(
        file_path,
        ["código", "número", "cantidad"],
        [["ARG", 1, 3], ["BRA", 1, 2], ["FRA", 1, 5]],
    )

    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_total == 3
    assert report.rows_applied == 3
    assert report.rows_skipped == 0
    assert _qty_for(db_conn, cards_by_key[("ARG", 1)]) == 3
    assert _qty_for(db_conn, cards_by_key[("BRA", 1)]) == 2
    assert _qty_for(db_conn, cards_by_key[("FRA", 1)]) == 5


def test_import_csv_same_behavior(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """CSV plano UTF-8: mismo behavior que xlsx."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "Messi")])
    file_path = tmp_path / "inv.csv"
    _write_csv(file_path, ["código", "número", "cantidad"], [["ARG", 1, 7]])

    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_applied == 1
    assert _qty_for(db_conn, cards[("ARG", 1)]) == 7


def test_import_csv_with_bom_is_handled(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """CSV con BOM UTF-8 (cómo Excel guarda en Windows) → header válido."""
    _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "Messi")])
    file_path = tmp_path / "inv_bom.csv"
    file_path.write_bytes(b"\xef\xbb\xbf" + b"c\xc3\xb3digo,n\xc3\xbamero,cantidad\nARG,1,4\n")
    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_applied == 1


def test_import_qty_zero_emits_warning_no_apply(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """qty=0 → warning, NO aplica. Pre-existente queda intacto."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "Messi")])
    # Pre-set qty=10 para verificar que no se toca.
    InventoryRepository(db_conn).adjust_quantity(cards[("ARG", 1)], 10)
    db_conn.commit()

    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(file_path, ["código", "número", "cantidad"], [["ARG", 1, 0]])

    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_applied == 0
    assert len(report.warnings) == 1
    assert report.warnings[0].code_id == "ARG"
    assert report.warnings[0].card_number == 1
    assert _qty_for(db_conn, cards[("ARG", 1)]) == 10  # intacto


def test_import_invalid_code_emits_error_does_not_abort(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """Código inexistente → 1 error; otra fila válida sí se aplica."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "Messi")])
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(
        file_path,
        ["código", "número", "cantidad"],
        [["ZZZ", 1, 5], ["ARG", 1, 3]],
    )
    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_applied == 1
    assert len(report.errors) == 1
    assert "ZZZ" in report.errors[0].message
    assert _qty_for(db_conn, cards[("ARG", 1)]) == 3


def test_import_invalid_number_emits_error(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """número 0, negativo, no entero → cada uno error."""
    _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "Messi")])
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(
        file_path,
        ["código", "número", "cantidad"],
        [["ARG", 0, 1], ["ARG", -3, 1], ["ARG", "abc", 1]],
    )
    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_applied == 0
    assert len(report.errors) == 3


def test_import_invalid_quantity_emits_error(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """cantidad negativa o no entera → error."""
    _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "Messi")])
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(
        file_path,
        ["código", "número", "cantidad"],
        [["ARG", 1, -2], ["ARG", 1, "xx"]],
    )
    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_applied == 0
    assert len(report.errors) == 2


def test_import_card_not_in_catalog_emits_error(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """Combinación válida sintácticamente pero no existe como card → error."""
    _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "Messi")])
    file_path = tmp_path / "inv.xlsx"
    # ARG-99 no existe en el catálogo.
    _write_xlsx(file_path, ["código", "número", "cantidad"], [["ARG", 99, 1]])
    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_applied == 0
    assert len(report.errors) == 1
    assert "no existe en el catálogo" in report.errors[0].message


def test_import_atomic_rollback_on_unhandled_exception(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Una excepción a mitad del bloque atómico → rollback total, ningún cambio."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "M"), ("BRA", 1, "N")])
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(
        file_path,
        ["código", "número", "cantidad"],
        [["ARG", 1, 2], ["BRA", 1, 5]],
    )

    # Hacer que el segundo `transactions.log` rompa.
    real_log = service._transactions.log
    call_count = {"n": 0}

    def flaky_log(*args: object, **kwargs: object) -> object:
        call_count["n"] += 1
        if call_count["n"] == 2:
            raise sqlite3.OperationalError("simulated")
        return real_log(*args, **kwargs)

    monkeypatch.setattr(service._transactions, "log", flaky_log)

    assert collection_with_codes.collection_id is not None
    with pytest.raises(sqlite3.OperationalError):
        service.import_inventory(
            file_path=file_path,
            collection_id=collection_with_codes.collection_id,
            mode="replace",
        )
    # Rollback: ninguna card quedó modificada.
    assert _qty_for(db_conn, cards[("ARG", 1)]) == 0
    assert _qty_for(db_conn, cards[("BRA", 1)]) == 0


def test_import_idempotent_in_replace(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """Mismo archivo dos veces en replace → mismo estado final."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "M")])
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(file_path, ["código", "número", "cantidad"], [["ARG", 1, 4]])
    assert collection_with_codes.collection_id is not None
    cid = collection_with_codes.collection_id
    service.import_inventory(file_path=file_path, collection_id=cid, mode="replace")
    service.import_inventory(file_path=file_path, collection_id=cid, mode="replace")
    assert _qty_for(db_conn, cards[("ARG", 1)]) == 4


def test_import_logs_transactions_with_shared_event_id(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """Tras un import de N filas con cambio efectivo, hay N transactions
    con el mismo exchange_event_id."""
    _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "M"), ("BRA", 1, "N")])
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(
        file_path,
        ["código", "número", "cantidad"],
        [["ARG", 1, 3], ["BRA", 1, 2]],
    )
    assert collection_with_codes.collection_id is not None
    service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    rows = db_conn.execute(
        "SELECT exchange_event_id, COUNT(*) FROM transactions " "GROUP BY exchange_event_id"
    ).fetchall()
    assert len(rows) == 1
    event_id, count = rows[0]
    assert event_id is not None
    assert count == 2


def test_import_replace_zero_delta_skips_transaction(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """En replace, qty del archivo == qty actual → no genera transaction."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "M")])
    InventoryRepository(db_conn).adjust_quantity(cards[("ARG", 1)], 5)
    db_conn.commit()
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(file_path, ["código", "número", "cantidad"], [["ARG", 1, 5]])
    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_applied == 1
    txn_count = db_conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0]
    assert txn_count == 0


def test_import_replace_negative_delta_logs_baja(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """En replace, qty del archivo < qty actual → BAJA por la diferencia."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "M")])
    InventoryRepository(db_conn).adjust_quantity(cards[("ARG", 1)], 8)
    db_conn.commit()
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(file_path, ["código", "número", "cantidad"], [["ARG", 1, 5]])
    assert collection_with_codes.collection_id is not None
    service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    txns = db_conn.execute("SELECT operation, quantity FROM transactions").fetchall()
    assert len(txns) == 1
    assert txns[0][0] == "baja"
    assert txns[0][1] == 3
    assert _qty_for(db_conn, cards[("ARG", 1)]) == 5


def test_import_strict_header_validation_raises(
    tmp_path: Path,
    service: InventoryImportService,
    collection_with_codes: Collection,
) -> None:
    """Header en inglés (no matchea) → ServiceError catastrófico."""
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(file_path, ["code", "number", "qty"], [["ARG", 1, 1]])
    assert collection_with_codes.collection_id is not None
    with pytest.raises(ServiceError, match="headers esperados"):
        service.import_inventory(
            file_path=file_path,
            collection_id=collection_with_codes.collection_id,
            mode="replace",
        )


def test_import_caps_errors_at_200(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """250 filas inválidas → errors[] truncado a 200."""
    _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "M")])
    file_path = tmp_path / "inv.xlsx"
    rows: list[list[object]] = [["ZZZ", 1, 1] for _ in range(250)]
    _write_xlsx(file_path, ["código", "número", "cantidad"], rows)
    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="replace",
    )
    assert report.rows_total == 250
    assert report.rows_skipped == 250  # contador real (no truncado)
    assert len(report.errors) == 200  # display truncado


# ---------------------------------------------------------------------
# import_inventory — modo add
# ---------------------------------------------------------------------


def test_import_add_mode_sums_to_existing(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """Pre-set qty=2; archivo qty=3 → final qty=5."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "M")])
    InventoryRepository(db_conn).adjust_quantity(cards[("ARG", 1)], 2)
    db_conn.commit()
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(file_path, ["código", "número", "cantidad"], [["ARG", 1, 3]])
    assert collection_with_codes.collection_id is not None
    report = service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="add",
    )
    assert report.rows_applied == 1
    assert _qty_for(db_conn, cards[("ARG", 1)]) == 5


def test_import_add_mode_starts_from_zero(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """Sin pre-existente: add con qty=4 → final qty=4 (alta por 4)."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "M")])
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(file_path, ["código", "número", "cantidad"], [["ARG", 1, 4]])
    assert collection_with_codes.collection_id is not None
    service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="add",
    )
    assert _qty_for(db_conn, cards[("ARG", 1)]) == 4


def test_import_add_mode_logs_alta_only(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_with_codes: Collection,
) -> None:
    """En add mode todas las transactions son ALTA (nunca BAJA)."""
    cards = _seed_cards(db_conn, collection_with_codes, [("ARG", 1, "M")])
    InventoryRepository(db_conn).adjust_quantity(cards[("ARG", 1)], 100)
    db_conn.commit()
    file_path = tmp_path / "inv.xlsx"
    _write_xlsx(file_path, ["código", "número", "cantidad"], [["ARG", 1, 1]])
    assert collection_with_codes.collection_id is not None
    service.import_inventory(
        file_path=file_path,
        collection_id=collection_with_codes.collection_id,
        mode="add",
    )
    ops = db_conn.execute("SELECT operation FROM transactions").fetchall()
    assert all(op[0] == "alta" for op in ops)


def test_import_unsupported_format_raises(
    tmp_path: Path,
    service: InventoryImportService,
    collection_with_codes: Collection,
) -> None:
    """Extensión .ods o similar → ServiceError catastrófico."""
    file_path = tmp_path / "inv.ods"
    file_path.write_text("dummy", encoding="utf-8")
    assert collection_with_codes.collection_id is not None
    with pytest.raises(ServiceError, match="formato no soportado"):
        service.import_inventory(
            file_path=file_path,
            collection_id=collection_with_codes.collection_id,
            mode="replace",
        )


# ---------------------------------------------------------------------
# Colección sin código (requires_code=False) — bug A1 de Sesión 5.5
# ---------------------------------------------------------------------


def test_import_inventory_collection_without_code(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_without_codes: Collection,
) -> None:
    """Sin columna `código` en el CSV; el lookup se resuelve por número.

    Las cards de la DB tienen `code_id="-"` (placeholder cualquier-cosa,
    NO necesariamente vacío). Pre-fix el importer usaba `code_id=""` en
    el lookup y fallaba con "('', N) no existe".
    """
    cards_repo = CardsRepository(db_conn)
    assert collection_without_codes.collection_id is not None
    cid = collection_without_codes.collection_id
    cards: dict[int, int] = {}
    for num, name in [(1, "Card uno"), (2, "Card dos"), (3, "Card tres")]:
        card = cards_repo.create(
            Card(
                card_id=None,
                collection_id=cid,
                code_id="-",  # placeholder, no es ""
                card_number=num,
                card_name=name,
            )
        )
        assert card.card_id is not None
        cards[num] = card.card_id
    db_conn.commit()

    file_path = tmp_path / "inv.csv"
    _write_csv(file_path, ["número", "cantidad"], [[1, 5], [2, 3], [3, 1]])

    report = service.import_inventory(
        file_path=file_path,
        collection_id=cid,
        mode="replace",
    )
    assert report.rows_total == 3
    assert report.rows_applied == 3
    assert report.rows_skipped == 0
    assert report.errors == []
    assert _qty_for(db_conn, cards[1]) == 5
    assert _qty_for(db_conn, cards[2]) == 3
    assert _qty_for(db_conn, cards[3]) == 1


def test_import_inventory_no_code_unknown_number_errors(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_without_codes: Collection,
) -> None:
    """Número que no existe en el catálogo → error claro, fila ignorada."""
    cards_repo = CardsRepository(db_conn)
    assert collection_without_codes.collection_id is not None
    cid = collection_without_codes.collection_id
    cards_repo.create(
        Card(card_id=None, collection_id=cid, code_id="-", card_number=1, card_name="A")
    )
    db_conn.commit()

    file_path = tmp_path / "inv.csv"
    _write_csv(file_path, ["número", "cantidad"], [[99, 1]])

    report = service.import_inventory(file_path=file_path, collection_id=cid, mode="replace")
    assert report.rows_applied == 0
    assert len(report.errors) == 1
    assert "99" in report.errors[0].message


def test_import_inventory_no_code_ambiguous_number_errors(
    tmp_path: Path,
    service: InventoryImportService,
    db_conn: sqlite3.Connection,
    collection_without_codes: Collection,
) -> None:
    """Misma colección sin código pero con dos cards de igual número
    (data inconsistente) → error claro, fila no importada."""
    cards_repo = CardsRepository(db_conn)
    assert collection_without_codes.collection_id is not None
    cid = collection_without_codes.collection_id
    # Dos cards con el mismo card_number=1 pero distinto code_id.
    cards_repo.create(
        Card(card_id=None, collection_id=cid, code_id="A", card_number=1, card_name="dup-A")
    )
    cards_repo.create(
        Card(card_id=None, collection_id=cid, code_id="B", card_number=1, card_name="dup-B")
    )
    db_conn.commit()

    file_path = tmp_path / "inv.csv"
    _write_csv(file_path, ["número", "cantidad"], [[1, 5]])

    report = service.import_inventory(file_path=file_path, collection_id=cid, mode="replace")
    assert report.rows_applied == 0
    assert len(report.errors) == 1
    assert "más de una vez" in report.errors[0].message
